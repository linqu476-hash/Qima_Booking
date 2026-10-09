"""Run management for the Streamlit front end. No Streamlit imports here."""
import io, json, os, re, shutil, signal, subprocess, sys, tempfile, time, zipfile
from datetime import datetime
from pathlib import Path
from uuid import uuid4

ROOT = Path(__file__).resolve().parent.parent
AUTOMATION = Path(os.getenv("QIMA_AUTOMATION_DIR") or ROOT / "automation")
RUNS = Path(os.getenv("QIMA_RUNS_DIR") or tempfile.gettempdir()) / "qima_runs"
RUNS.mkdir(parents=True, exist_ok=True)

# Steps are inferred from the screenshot names main.py already writes.
STEPS = [
    ("Login", ["login_page", "login_filled", "after_login"]),
    ("Booking", ["all_services", "booking_form", "booking_ref_done"]),
    ("Supplier", ["supplier_popup", "supplier_row_selected", "supplier_done"]),
    ("Date and Excel", ["date_time_done", "excel_uploaded"]),
    ("Products", ["after_next", "product_info", "po_all_done"]),
    ("Inspection", ["inspection_details"]),
]
FAIL_SHOT_RE = re.compile(r"error|fail|missing|not_found|blocked|no_po", re.I)
BOOKING_RE = re.compile(r"===== booking (\d+)/(\d+)\s+ref=(\S+)")
SHOT_RE = re.compile(r"screenshot -> (\S+\.png)")
RESULT_RE = re.compile(r"INFO (\S+)\s+->\s+(DONE|FAILED|DRYRUN_OK)\s*$", re.M)


def now() -> str:
    return datetime.now().isoformat(timespec="seconds")


def load_meta(rd: Path) -> dict:
    return json.loads((rd / "run.json").read_text())


def save_meta(rd: Path, m: dict) -> None:
    (rd / "run.json").write_text(json.dumps(m, indent=2))


def list_runs() -> list:
    out = []
    for p in sorted(RUNS.iterdir(), reverse=True):
        if (p / "run.json").exists():
            out.append(load_meta(p))
    return out


def env_file_text(env: dict) -> str:
    return "".join("{}='{}'\n".format(k, str(v).replace("\\", "\\\\").replace("'", "\\'")) for k, v in env.items())


def create_run(booking, booking_list, pdfs, dry_run: bool, base: Path = None, ti=None) -> Path:
    """booking/booking_list are bytes or None; pdfs is a list of (name, bytes).
    Empty slots reuse the previous run's files, including the Status column the script wrote,
    so rows already DONE are skipped."""
    rid = time.strftime("%Y%m%d-%H%M%S") + "-" + uuid4().hex[:4]
    rd = RUNS / rid
    (rd / "input" / "po_pdf").mkdir(parents=True)
    prev = base or next((p for p in sorted(RUNS.iterdir(), reverse=True)
                         if p.name != rid and (p / "input").is_dir()), None)
    if prev and (prev / "input").is_dir():
        shutil.copytree(prev / "input", rd / "input", dirs_exist_ok=True)
    if booking:
        (rd / "input" / "booking.xlsx").write_bytes(booking)
    if booking_list:
        (rd / "input" / "booking_list.xlsx").write_bytes(booking_list)
    if ti:
        (rd / "input" / "technical_info.xlsx").write_bytes(ti)
    if pdfs:
        shutil.rmtree(rd / "input" / "po_pdf", ignore_errors=True)
        (rd / "input" / "po_pdf").mkdir()
        for name, data in pdfs:
            (rd / "input" / "po_pdf" / Path(name).name).write_bytes(data)
    if not any((rd / "input").rglob("*.*")):
        shutil.rmtree(rd)
        raise ValueError("Upload at least one input file.")
    shutil.copytree(AUTOMATION, rd, dirs_exist_ok=True, ignore=shutil.ignore_patterns(
        "input", "logs", ".env", "__pycache__", "*.rar", ".git", "venv", ".venv"))
    (rd / "logs").mkdir(exist_ok=True)
    save_meta(rd, {"id": rid, "created": now(), "status": "queued", "dry_run": dry_run, "note": "",
                   "rc": None, "results": [], "has_errors": False, "started_ts": None, "finished": None,
                   "inherit_list": not booking_list})  # no new list uploaded: take the latest finished run's Status column
    return rd


def start(rd: Path, env_extra: dict, dry_run: bool) -> subprocess.Popen:
    env = {**os.environ, **env_extra, "HEADLESS": "true", "DRY_RUN": "true" if dry_run else "false",
           "PYTHONUNBUFFERED": "1", "PYTHONIOENCODING": "utf-8"}
    (rd / ".env").write_text(env_file_text(env_extra))
    with open(rd / "run.log", "ab") as log:
        proc = subprocess.Popen([sys.executable, "-u", "main.py"], cwd=str(rd), env=env, stdout=log,
                                stderr=subprocess.STDOUT, start_new_session=(os.name != "nt"))
    m = load_meta(rd)
    m.update(status="running", started_ts=time.time())
    save_meta(rd, m)
    return proc


def terminate(proc) -> None:
    try:
        if os.name == "nt":
            subprocess.run(["taskkill", "/T", "/F", "/PID", str(proc.pid)], capture_output=True)
        else:
            os.killpg(os.getpgid(proc.pid), signal.SIGTERM)
        proc.wait(timeout=8)
    except subprocess.TimeoutExpired:
        if os.name != "nt":
            os.killpg(os.getpgid(proc.pid), signal.SIGKILL)
    except ProcessLookupError:
        pass


def stop(rd: Path, proc) -> None:
    m = load_meta(rd)
    if m["status"] == "queued":
        m.update(status="stopped", finished=now(), note="Cancelled before it started")
        save_meta(rd, m)
        return
    m["stopping"] = True
    save_meta(rd, m)
    if proc is not None and proc.poll() is None:
        terminate(proc)


def log_info(rd: Path):
    f = rd / "run.log"
    text = f.read_text("utf-8", "replace") if f.exists() else ""
    cur, shots = None, []
    for line in text.splitlines():
        mm = BOOKING_RE.search(line)
        if mm:
            cur = {"n": int(mm[1]), "total": int(mm[2]), "ref": mm[3]}
            shots = [x for x in shots if any(k in x for k in STEPS[0][1])]  # keep login shots
            continue
        mm = SHOT_RE.search(line)
        if mm:
            shots.append(mm[1])
    return cur, shots, [{"ref": a, "status": b} for a, b in RESULT_RE.findall(text)], text


def finalize(rd: Path, m: dict, rc, timed_out: bool = False, timeout_min: int = 0) -> dict:
    _, _, results, text = log_info(rd)
    lines = [ln.strip() for ln in text.splitlines() if ln.strip()]
    # prefer the final "SomeError: message" line over trailing log noise
    last = next((ln for ln in reversed(lines) if re.match(r"^[\w.]*(Error|Exception)\b", ln)), lines[-1] if lines else "")
    n_failed = sum(r["status"] == "FAILED" for r in results)
    m.update(rc=rc, finished=now(), results=results,
             has_errors=bool(n_failed) or any(FAIL_SHOT_RE.search(p.name) for p in (rd / "logs").glob("*.png")))
    if m.get("stopping"):
        m.update(status="stopped", note="Stopped from the dashboard")
    elif timed_out:
        m.update(status="failed", note=f"Timed out after {timeout_min} minutes")
    elif rc != 0 and "Nothing to do" in last:  # every row already DONE
        m.update(status="done", has_errors=False, note=last[:200])
    elif rc != 0:
        m.update(status="failed", note=(last or f"Script exited with code {rc}")[:200])
    elif results and n_failed == len(results):
        m.update(status="failed", note="Every booking failed. See the Bookings table and screenshots.")
    elif n_failed:
        m.update(status="done", note=f"{n_failed} of {len(results)} bookings failed. Retry runs only the failed rows.")
    else:
        m.update(status="done", note="Finished, but error screenshots were saved. Check the gallery." if m["has_errors"] else "")
    save_meta(rd, m)
    return m


def refresh_status(rd: Path, proc, timeout_min: int) -> dict:
    m = load_meta(rd)
    if m["status"] != "running":
        return m
    if proc is not None and proc.poll() is None:
        if time.time() - m["started_ts"] > timeout_min * 60:
            terminate(proc)
            return finalize(rd, m, proc.returncode, timed_out=True, timeout_min=timeout_min)
        return m
    if proc is None:  # the app restarted, the script died with it
        m.update(status="failed", finished=now(), note="Interrupted: the app restarted while the run was in progress")
        save_meta(rd, m)
        return m
    return finalize(rd, m, proc.returncode)


def progress(names, status: str) -> list:
    names = [n for n in names if not FAIL_SHOT_RE.search(n)]
    reached = -1
    for i, (_, keys) in enumerate(STEPS):
        if any(k in n for n in names for k in keys):
            reached = i
    out = []
    for i, (label, _) in enumerate(STEPS):
        at = max(reached, 0)
        if status == "done":
            s = "done"
        elif status in ("created", "queued"):
            s = "pending"
        elif status in ("failed", "stopped"):
            s = "done" if i < at else (status if i == at else "pending")
        else:
            s = "done" if i < at else ("active" if i == at else "pending")
        out.append((label, s))
    return out


def read_bookings(rd: Path) -> list:
    try:
        from openpyxl import load_workbook
        rows = list(load_workbook(rd / "input" / "booking_list.xlsx", read_only=True, data_only=True)["Bookings"].iter_rows(values_only=True))
        head = [str(c or "").strip().lower() for c in rows[0]]

        def g(r, k):
            v = r[head.index(k)] if k in head else ""
            v = int(v) if isinstance(v, float) and v.is_integer() else v
            return "" if v is None else str(v).strip()
        return [{"Booking ref": g(r, "booking ref no."), "Supplier": g(r, "supplier"),
                 "Status": g(r, "status").upper() or "Pending", "Message": g(r, "message")}
                for r in rows[1:] if g(r, "booking ref no.")]
    except Exception:
        return []


def make_zip(rd: Path) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for sub in ("input", "logs"):
            for p in (rd / sub).rglob("*"):
                if p.is_file():
                    z.write(p, p.relative_to(rd).as_posix())
        if (rd / "run.log").exists():
            z.write(rd / "run.log", "run.log")
    return buf.getvalue()


# ---------------------------------------------------------------- run queue
def _all_meta() -> list:
    return [load_meta(p) for p in sorted(RUNS.iterdir()) if (p / "run.json").exists()]  # oldest first


def queue_ids() -> list:
    return [m["id"] for m in _all_meta() if m["status"] == "queued"]


def _start_queued(rd: Path, m: dict, env: dict):
    if m.get("inherit_list"):  # use the Status column the previous finished run wrote, so DONE rows are never booked twice
        for p in sorted(RUNS.iterdir(), reverse=True):
            if p.name >= rd.name or not (p / "run.json").exists():
                continue
            if load_meta(p)["status"] in ("queued", "running"):
                continue
            src = p / "input" / "booking_list.xlsx"
            if src.exists():
                shutil.copy2(src, rd / "input" / "booking_list.xlsx")
            break
    return start(rd, env, m.get("dry_run", False))


def tick(rt: dict, cfg: dict) -> None:
    """One pass of the worker: finish ended runs, then start queued runs while slots are free."""
    with rt["lock"]:
        procs = rt["procs"]
        for rid, proc in list(procs.items()):
            if refresh_status(RUNS / rid, proc, cfg["timeout"])["status"] != "running":
                procs.pop(rid, None)
        metas = _all_meta()
        for m in metas:  # marked running but no process: the app restarted
            if m["status"] == "running" and m["id"] not in procs:
                refresh_status(RUNS / m["id"], None, cfg["timeout"])
        for m in metas:
            if len(procs) >= cfg.get("max_parallel", 1):
                break
            if m["status"] != "queued":
                continue
            rd = RUNS / m["id"]
            try:
                procs[m["id"]] = _start_queued(rd, m, cfg["env"])
            except Exception as e:
                m.update(status="failed", finished=now(), note=f"Could not start: {e}")
                save_meta(rd, m)
