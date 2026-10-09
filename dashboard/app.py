"""QIMA RPA dashboard: runs automation/main.py as managed jobs.

Each run gets its own workspace (data/runs/<id>/) containing a copy of the
automation folder, the uploaded input/ files, a generated .env, and logs/.
"""
import asyncio, hashlib, hmac, io, json, os, re, shutil, signal, smtplib, sys, time, urllib.request, zipfile
from contextlib import asynccontextmanager
from datetime import datetime
from email.message import EmailMessage
from pathlib import Path
from typing import List, Optional
from uuid import uuid4

from fastapi import Depends, FastAPI, File, Form, HTTPException, Request, Response, UploadFile
from fastapi.responses import FileResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles

BASE = Path(__file__).resolve().parent
DATA = Path(os.getenv("DATA_DIR", BASE / "data")).resolve()
RUNS = DATA / "runs"
RUNS.mkdir(parents=True, exist_ok=True)
AUTOMATION = Path(os.getenv("AUTOMATION_DIR", BASE.parent / "automation")).resolve()
SETTINGS_FILE = DATA / "settings.json"
PASSWORD = os.getenv("DASH_PASSWORD", "")
SECRET = os.getenv("SECRET_KEY") or hashlib.sha256(("qima:" + PASSWORD).encode()).hexdigest()
PUBLIC_URL = os.getenv("PUBLIC_URL", "").rstrip("/")
if not PASSWORD:
    sys.exit("Set DASH_PASSWORD before starting the dashboard.")

MASK = "********"
SECRET_KEY_RE = re.compile(r"pass|pwd|secret|token|key", re.I)
RUN_ID_RE = re.compile(r"^\d{8}-\d{6}-[0-9a-f]{4}$")

# Steps are inferred from the screenshot names your script already writes to logs/.
STEPS = [
    ("login", "Login", ["login_page", "login_filled", "after_login"]),
    ("booking", "Booking", ["all_services", "booking_form", "booking_ref_done"]),
    ("supplier", "Supplier", ["supplier_popup", "supplier_row_selected", "supplier_done"]),
    ("upload", "Date and Excel", ["date_time_done", "excel_uploaded"]),
    ("products", "Products", ["after_next", "product_info", "po_all_done"]),
    ("inspection", "Inspection", ["inspection_details"]),
]

FAIL_SHOT_RE = re.compile(r"error|fail|missing|not_found|blocked|no_po", re.I)
BOOKING_RE = re.compile(r"===== booking (\d+)/(\d+)\s+ref=(\S+)")
SHOT_RE = re.compile(r"screenshot -> (\S+\.png)")
RESULT_RE = re.compile(r"INFO (\S+)\s+->\s+(DONE|FAILED|DRYRUN_OK)\s*$", re.M)

DEFAULTS = {
    "env": {},
    "headless": True,
    "dry_run": False,
    "timeout_min": 60,
    "alert_on": "problems",  # problems | always | never
    "slack_webhook": "",
    "email": {"enabled": False, "host": "", "port": 587, "user": "", "password": "", "from": "", "to": ""},
    "schedule": {"enabled": False, "time": "09:00", "days": [0, 1, 2, 3, 4], "last_fired": ""},
}

QUEUE: "asyncio.Queue[str]" = asyncio.Queue()
PROCS: dict = {}
STOPPING: set = set()


def now() -> str:
    return datetime.now().isoformat(timespec="seconds")


# ---------------------------------------------------------------- settings
def load_settings() -> dict:
    s = json.loads(json.dumps(DEFAULTS))
    if SETTINGS_FILE.exists():
        saved = json.loads(SETTINGS_FILE.read_text())
        for k, v in saved.items():
            if isinstance(v, dict) and isinstance(s.get(k), dict) and k != "env":
                s[k].update(v)
            else:
                s[k] = v
    return s


def save_settings(s: dict) -> None:
    SETTINGS_FILE.write_text(json.dumps(s, indent=2))
    os.chmod(SETTINGS_FILE, 0o600)


def public_settings(s: dict) -> dict:
    p = json.loads(json.dumps(s))
    p["env"] = {k: (MASK if SECRET_KEY_RE.search(k) and v else v) for k, v in p["env"].items()}
    if p["slack_webhook"]:
        p["slack_webhook"] = MASK
    if p["email"]["password"]:
        p["email"]["password"] = MASK
    return p


# -------------------------------------------------------------------- auth
def _sign(ts: str) -> str:
    return hmac.new(SECRET.encode(), ts.encode(), "sha256").hexdigest()


def make_token() -> str:
    ts = str(int(time.time()))
    return f"{ts}.{_sign(ts)}"


def token_ok(tok: str) -> bool:
    try:
        ts, sig = tok.split(".")
        return hmac.compare_digest(sig, _sign(ts)) and time.time() - int(ts) < 7 * 86400
    except Exception:
        return False


def auth(request: Request) -> None:
    if not token_ok(request.cookies.get("session", "")):
        raise HTTPException(401, "Sign in required")


# -------------------------------------------------------------------- runs
def run_dir(rid: str) -> Path:
    if not RUN_ID_RE.match(rid) or not (RUNS / rid).is_dir():
        raise HTTPException(404, "Run not found")
    return RUNS / rid


def load_run(rid: str) -> dict:
    return json.loads((run_dir(rid) / "run.json").read_text())


def save_run(m: dict) -> None:
    (RUNS / m["id"] / "run.json").write_text(json.dumps(m, indent=2))


def all_runs() -> List[dict]:
    out = []
    for p in sorted(RUNS.iterdir(), reverse=True):
        f = p / "run.json"
        if f.exists():
            out.append(json.loads(f.read_text()))
    return out


def create_run(booking: Optional[UploadFile], booking_list: Optional[UploadFile],
               pdfs: List[UploadFile], base: Optional[str], source: str) -> dict:
    pdfs = [p for p in pdfs if p.filename]
    runs = all_runs()
    if not (booking and booking.filename and booking_list and booking_list.filename and pdfs):
        # some files come from a previous run, which must be finished so its Status column is final
        if any(r["status"] in ("queued", "running") for r in runs):
            raise HTTPException(409, "A run is still in progress. Wait for it to finish, or upload all three files.")
    rid = time.strftime("%Y%m%d-%H%M%S") + "-" + uuid4().hex[:4]
    d = RUNS / rid
    (d / "input" / "po_pdf").mkdir(parents=True)
    (d / "logs").mkdir()
    # Reuse the previous run's input files, including its booking_list.xlsx with the Status column
    # the script wrote, so rows already marked DONE are skipped and are never booked twice.
    prev = RUNS / base if base and RUN_ID_RE.match(base) and (RUNS / base).is_dir() else (RUNS / runs[0]["id"] if runs else None)
    if prev and (prev / "input").is_dir():
        shutil.copytree(prev / "input", d / "input", dirs_exist_ok=True)

    def put(up: UploadFile, dest: Path):
        dest.parent.mkdir(parents=True, exist_ok=True)
        with open(dest, "wb") as f:
            shutil.copyfileobj(up.file, f)

    if booking and booking.filename:
        put(booking, d / "input" / "booking.xlsx")
    if booking_list and booking_list.filename:
        put(booking_list, d / "input" / "booking_list.xlsx")
    if pdfs:
        shutil.rmtree(d / "input" / "po_pdf", ignore_errors=True)
        for p in pdfs:
            put(p, d / "input" / "po_pdf" / Path(p.filename).name)
    if not any((d / "input").rglob("*.*")):
        shutil.rmtree(d)
        raise HTTPException(400, "Upload at least one input file.")
    m = {"id": rid, "status": "queued", "created": now(), "started": None, "finished": None,
         "rc": None, "source": source, "has_errors": False, "note": "", "baseline": [], "results": [],
         "inputs": sorted(str(p.relative_to(d / "input")) for p in (d / "input").rglob("*") if p.is_file())}
    save_run(m)
    QUEUE.put_nowait(rid)
    return m


def log_info(d: Path):
    """Parse the script's own log lines: current booking, screenshots of that booking, final results."""
    f = d / "run.log"
    text = f.read_text("utf-8", "replace") if f.exists() else ""
    cur, shots = None, []
    for line in text.splitlines():
        mm = BOOKING_RE.search(line)
        if mm:
            cur = {"n": int(mm[1]), "total": int(mm[2]), "ref": mm[3]}
            shots = [x for x in shots if any(k in x for k in STEPS[0][2])]  # keep the login shots
            continue
        mm = SHOT_RE.search(line)
        if mm:
            shots.append(mm[1])
    return cur, shots, [{"ref": a, "status": b} for a, b in RESULT_RE.findall(text)], text


def read_bookings(d: Path) -> List[dict]:
    """Rows of booking_list.xlsx (the script writes Status and Message back into it)."""
    try:
        from openpyxl import load_workbook
        rows = list(load_workbook(d / "input" / "booking_list.xlsx", read_only=True, data_only=True)["Bookings"].iter_rows(values_only=True))
        head = [str(c or "").strip().lower() for c in rows[0]]
        def g(r, k):
            v = r[head.index(k)] if k in head else ""
            v = int(v) if isinstance(v, float) and v.is_integer() else v
            return "" if v is None else str(v).strip()
        return [{"ref": g(r, "booking ref no."), "supplier": g(r, "supplier"), "status": g(r, "status").upper(),
                 "message": g(r, "message")} for r in rows[1:] if g(r, "booking ref no.")]
    except Exception:
        return []


def progress(names: List[str], status: str) -> List[dict]:
    names = [n for n in names if not FAIL_SHOT_RE.search(n)]
    reached = -1
    for i, (_, _, keys) in enumerate(STEPS):
        if any(k in n for n in names for k in keys):
            reached = i
    out = []
    for i, (key, label, _) in enumerate(STEPS):
        if status == "done":
            state = "done"
        elif status == "queued":
            state = "pending"
        elif status in ("failed", "stopped"):
            at = max(reached, 0)
            state = "done" if i < at else (("failed" if status == "failed" else "stopped") if i == at else "pending")
        else:
            at = max(reached, 0)
            state = "done" if i < at else ("active" if i == at else "pending")
        out.append({"key": key, "label": label, "state": state})
    return out


def run_files(d: Path, m: dict) -> List[dict]:
    base = set(m.get("baseline", []))
    files = []
    for p in sorted(d.rglob("*")):
        if not p.is_file():
            continue
        rel = p.relative_to(d).as_posix()
        top = rel.split("/")[0]
        if top in ("input_original", "logs", "__pycache__") or rel in base or rel in (".env", "run.json", "run.log"):
            continue
        files.append({"path": rel, "size": p.stat().st_size, "kind": "input" if top == "input" else "output"})
    return files


async def terminate(proc) -> None:
    if os.name == "nt":  # Windows: kill the script and the browser processes it started
        k = await asyncio.create_subprocess_exec("taskkill", "/T", "/F", "/PID", str(proc.pid),
                                                 stdout=asyncio.subprocess.DEVNULL, stderr=asyncio.subprocess.DEVNULL)
        await k.wait()
        return
    try:
        os.killpg(os.getpgid(proc.pid), signal.SIGTERM)
        await asyncio.wait_for(proc.wait(), 8)
    except asyncio.TimeoutError:
        os.killpg(os.getpgid(proc.pid), signal.SIGKILL)
    except ProcessLookupError:
        pass


def env_file_text(env: dict) -> str:
    lines = []
    for k, v in env.items():
        v = str(v).replace("\\", "\\\\").replace("'", "\\'")
        lines.append(f"{k}='{v}'\n")
    return "".join(lines)


async def execute(rid: str) -> None:
    s = load_settings()
    d = RUNS / rid
    m = load_run(rid)
    m.update(status="running", started=now())
    save_run(m)
    if not (AUTOMATION / "main.py").exists():
        m.update(status="failed", finished=now(), note=f"{AUTOMATION}/main.py not found")
        save_run(m)
        return await notify(m)
    shutil.copytree(AUTOMATION, d, dirs_exist_ok=True, ignore=shutil.ignore_patterns(
        "input", "logs", ".env", "__pycache__", "*.rar", ".git", "venv", ".venv"))
    (d / "logs").mkdir(exist_ok=True)
    (d / ".env").write_text(env_file_text(s["env"]))
    m["baseline"] = sorted(p.relative_to(d).as_posix() for p in d.rglob("*")
                           if p.is_file() and p.relative_to(d).parts[0] not in ("input", "input_original", "logs"))
    save_run(m)

    env = {**os.environ, **{k: str(v) for k, v in s["env"].items()},
           "HEADLESS": "true" if s["headless"] else "false", "DRY_RUN": "true" if s["dry_run"] else "false",
           "PYTHONUNBUFFERED": "1", "PYTHONIOENCODING": "utf-8"}
    cmd = [sys.executable, "-u", "main.py"]
    # Only needed for "Visible" mode: a virtual display so a headed browser works on a server.
    if not s["headless"] and shutil.which("xvfb-run") and shutil.which("xauth"):
        cmd = ["xvfb-run", "-a"] + cmd
    timed_out = False
    with open(d / "run.log", "ab") as log:
        proc = await asyncio.create_subprocess_exec(*cmd, cwd=str(d), env=env, stdout=log,
                                                    stderr=asyncio.subprocess.STDOUT, start_new_session=(os.name != "nt"))
        PROCS[rid] = proc
        try:
            rc = await asyncio.wait_for(proc.wait(), int(s["timeout_min"]) * 60)
        except asyncio.TimeoutError:
            timed_out = True
            await terminate(proc)
            rc = proc.returncode
        finally:
            PROCS.pop(rid, None)
    m = load_run(rid)
    _, _, results, text = log_info(d)
    last = next((ln.strip() for ln in reversed(text.splitlines()) if ln.strip()), "")
    n_failed = sum(r["status"] == "FAILED" for r in results)
    m.update(rc=rc, finished=now(), results=results,
             has_errors=bool(n_failed) or any(FAIL_SHOT_RE.search(p.name) for p in (d / "logs").glob("*.png")))
    if rid in STOPPING:
        STOPPING.discard(rid)
        m.update(status="stopped", note="Stopped from the dashboard")
    elif timed_out:
        m.update(status="failed", note=f"Timed out after {s['timeout_min']} minutes")
    elif rc != 0 and "Nothing to do" in last:  # every row already DONE: not a failure
        m.update(status="done", has_errors=False, note=last[:200])
    elif rc != 0:
        m.update(status="failed", note=(last or f"Script exited with code {rc}")[:200])
    elif results and n_failed == len(results):
        m.update(status="failed", note="Every booking failed. See the Bookings table and screenshots.")
    elif n_failed:
        m.update(status="done", note=f"{n_failed} of {len(results)} bookings failed. Retry runs only the failed rows.")
    else:
        m.update(status="done", note="Finished, but error screenshots were saved. Check the gallery." if m["has_errors"] else "")
    save_run(m)
    await notify(m)


async def worker() -> None:
    while True:
        rid = await QUEUE.get()
        try:
            if load_run(rid)["status"] == "queued":
                await execute(rid)
        except Exception as e:  # keep the queue alive whatever happens
            print("run failed to start:", rid, e, file=sys.stderr)
            try:
                m = load_run(rid)
                m.update(status="failed", finished=now(), note=f"Dashboard error: {e}")
                save_run(m)
            except Exception:
                pass
        finally:
            QUEUE.task_done()


# ----------------------------------------------------------- alerts + cron
def send_alerts(s: dict, subject: str, text: str) -> List[str]:
    errors = []
    if s["slack_webhook"]:
        try:
            req = urllib.request.Request(s["slack_webhook"], json.dumps({"text": f"*{subject}*\n{text}"}).encode(),
                                         {"Content-Type": "application/json"})
            urllib.request.urlopen(req, timeout=15).read()
        except Exception as e:
            errors.append(f"Slack: {e}")
    e = s["email"]
    if e["enabled"] and e["host"] and e["to"]:
        try:
            msg = EmailMessage()
            msg["Subject"], msg["From"], msg["To"] = subject, e["from"] or e["user"], e["to"]
            msg.set_content(text)
            with smtplib.SMTP(e["host"], int(e["port"]), timeout=20) as smtp:
                smtp.starttls()
                if e["user"]:
                    smtp.login(e["user"], e["password"])
                smtp.send_message(msg)
        except Exception as ex:
            errors.append(f"Email: {ex}")
    return errors


async def notify(m: dict) -> None:
    s = load_settings()
    problem = m["status"] == "failed" or m.get("has_errors")
    if m["status"] == "stopped" or s["alert_on"] == "never" or (s["alert_on"] == "problems" and not problem):
        return
    label = "needs attention" if problem else "finished"
    link = f"\n{PUBLIC_URL}/#run={m['id']}" if PUBLIC_URL else ""
    text = f"Run {m['id']} {label} (status: {m['status']}). {m.get('note', '')}{link}"
    errs = await asyncio.to_thread(send_alerts, s, f"QIMA RPA run {label}", text)
    for e in errs:
        print("alert error:", e, file=sys.stderr)


async def scheduler() -> None:
    while True:
        await asyncio.sleep(20)
        try:
            s = load_settings()
            sc = s["schedule"]
            t = datetime.now()
            key = t.strftime("%Y-%m-%d ") + sc["time"]
            busy = any(m["status"] in ("queued", "running") for m in all_runs()[:5])
            if sc["enabled"] and not busy and t.strftime("%H:%M") == sc["time"] and t.weekday() in sc["days"] and sc["last_fired"] != key:
                s["schedule"]["last_fired"] = key
                save_settings(s)
                create_run(None, None, [], None, "schedule")
        except Exception as e:
            print("scheduler:", e, file=sys.stderr)


@asynccontextmanager
async def lifespan(_: FastAPI):
    for m in all_runs():  # runs left over from a restart can't resume
        if m["status"] in ("queued", "running"):
            m.update(status="failed", finished=now(), note="Interrupted by a dashboard restart")
            save_run(m)
    tasks = [asyncio.create_task(worker()), asyncio.create_task(scheduler())]
    yield
    for t in tasks:
        t.cancel()


app = FastAPI(lifespan=lifespan)


# ------------------------------------------------------------------- routes
@app.post("/api/login")
async def login(request: Request, response: Response):
    body = await request.json()
    if not hmac.compare_digest(str(body.get("password", "")), PASSWORD):
        await asyncio.sleep(1)
        raise HTTPException(401, "Wrong password")
    response.set_cookie("session", make_token(), httponly=True, samesite="lax",
                        secure=os.getenv("COOKIE_SECURE") == "1", max_age=7 * 86400)
    return {"ok": True}


@app.post("/api/logout")
def logout(response: Response):
    response.delete_cookie("session")
    return {"ok": True}


@app.get("/api/me", dependencies=[Depends(auth)])
def me():
    return {"ok": True, "queue": QUEUE.qsize(), "script": (AUTOMATION / "main.py").exists()}


@app.get("/api/runs", dependencies=[Depends(auth)])
def list_runs():
    return [{k: v for k, v in m.items() if k not in ("baseline", "inputs")} for m in all_runs()[:100]]


@app.post("/api/runs", dependencies=[Depends(auth)])
async def new_run(booking: Optional[UploadFile] = File(None), booking_list: Optional[UploadFile] = File(None),
            po_pdfs: List[UploadFile] = File(default=[]), base: Optional[str] = Form(None)):
    return create_run(booking, booking_list, po_pdfs, base, "manual")


@app.get("/api/runs/{rid}", dependencies=[Depends(auth)])
def run_detail(rid: str, offset: int = 0):
    d = run_dir(rid)
    m = load_run(rid)
    text, new_offset = "", offset
    if (d / "run.log").exists():
        with open(d / "run.log", "rb") as f:
            f.seek(offset)
            chunk = f.read(200_000)
            new_offset = offset + len(chunk)
            text = chunk.decode("utf-8", "replace")
    shots = [{"name": p.name, "error": bool(FAIL_SHOT_RE.search(p.name))} for p in sorted((d / "logs").glob("*.png"))]
    cur, step_shots, _, _ = log_info(d)
    pos = list(QUEUE._queue).index(rid) + 1 if rid in list(QUEUE._queue) else 0  # type: ignore[attr-defined]
    return {"run": {k: v for k, v in m.items() if k != "baseline"}, "steps": progress(step_shots, m["status"]), "current": cur, "bookings": read_bookings(d),
            "log": text, "offset": new_offset, "shots": shots, "files": run_files(d, m), "queue_position": pos}


@app.post("/api/runs/{rid}/stop", dependencies=[Depends(auth)])
async def stop_run(rid: str):
    m = load_run(rid)
    if m["status"] == "queued":
        m.update(status="stopped", finished=now(), note="Cancelled before it started")
        save_run(m)
    elif m["status"] == "running" and rid in PROCS:
        STOPPING.add(rid)
        asyncio.create_task(terminate(PROCS[rid]))
    return {"ok": True}


@app.post("/api/runs/{rid}/retry", dependencies=[Depends(auth)])
async def retry_run(rid: str):
    run_dir(rid)
    return create_run(None, None, [], rid, "retry")


@app.delete("/api/runs/{rid}", dependencies=[Depends(auth)])
def delete_run(rid: str):
    d = run_dir(rid)
    if load_run(rid)["status"] in ("queued", "running"):
        raise HTTPException(409, "Stop the run before deleting it")
    shutil.rmtree(d)
    return {"ok": True}


def safe_file(rid: str, path: str) -> Path:
    d = run_dir(rid)
    target = (d / path).resolve()
    if not target.is_relative_to(d) or not target.is_file() or target.name in (".env", "run.json") \
            or target.relative_to(d).parts[0] == "input_original":
        raise HTTPException(404, "File not found")
    return target


@app.get("/api/runs/{rid}/file/{path:path}", dependencies=[Depends(auth)])
def get_file(rid: str, path: str, download: int = 0):
    f = safe_file(rid, path)
    return FileResponse(f, filename=f.name if download else None,
                        content_disposition_type="attachment" if download else "inline")


@app.get("/api/runs/{rid}/zip", dependencies=[Depends(auth)])
def get_zip(rid: str):
    d = run_dir(rid)
    m = load_run(rid)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for f in run_files(d, m):
            z.write(d / f["path"], f["path"])
        for p in (d / "logs").glob("*"):
            z.write(p, f"logs/{p.name}")
        if (d / "run.log").exists():
            z.write(d / "run.log", "run.log")
    buf.seek(0)
    return StreamingResponse(buf, media_type="application/zip",
                             headers={"Content-Disposition": f'attachment; filename="run-{rid}.zip"'})


@app.get("/api/settings", dependencies=[Depends(auth)])
def get_settings():
    return public_settings(load_settings())


@app.put("/api/settings", dependencies=[Depends(auth)])
async def put_settings(request: Request):
    new, old = await request.json(), load_settings()
    s = json.loads(json.dumps(DEFAULTS))
    s["env"] = {k.strip(): (old["env"].get(k, "") if v == MASK else str(v))
                for k, v in new.get("env", {}).items() if k.strip() and re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]*", k.strip())}
    s["headless"] = bool(new.get("headless", True))
    s["dry_run"] = bool(new.get("dry_run", False))
    s["timeout_min"] = max(1, int(new.get("timeout_min", 60)))
    s["alert_on"] = new.get("alert_on", "problems") if new.get("alert_on") in ("problems", "always", "never") else "problems"
    s["slack_webhook"] = old["slack_webhook"] if new.get("slack_webhook") == MASK else new.get("slack_webhook", "").strip()
    e = {**DEFAULTS["email"], **new.get("email", {})}
    if e["password"] == MASK:
        e["password"] = old["email"]["password"]
    s["email"] = {**e, "enabled": bool(e["enabled"]), "port": int(e["port"] or 587)}
    sc = new.get("schedule", {})
    s["schedule"] = {"enabled": bool(sc.get("enabled")), "time": sc.get("time", "09:00"),
                     "days": [int(x) for x in sc.get("days", [0, 1, 2, 3, 4])], "last_fired": old["schedule"]["last_fired"]}
    save_settings(s)
    return public_settings(s)


@app.post("/api/settings/test-alert", dependencies=[Depends(auth)])
async def test_alert():
    errs = await asyncio.to_thread(send_alerts, load_settings(), "QIMA RPA test alert", "Alerts are working.")
    return {"ok": not errs, "errors": errs}


app.mount("/", StaticFiles(directory=BASE / "static", html=True), name="static")
